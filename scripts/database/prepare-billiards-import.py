#!/usr/bin/env python3
"""Prepare a reviewable Amap-file import. No database writes and no legacy-customer matching."""
import argparse
from collections import Counter
import csv
import hashlib
import io
import json
import math
from pathlib import Path
import re
import unicodedata
from zipfile import ZipFile

FULL_SHEETS = ('全部精准门店', '名称近似全部门店')
REQUIRED = {'门店名称', '地址', '经纬度', '高德 POI ID', '行政区划代码'}
VENUE = re.compile(r'台球|桌球|撞球|斯诺克|billiards?|snooker', re.I)
SUPPLIER = re.compile(r'用品|器材|球杆|制造|工厂|配件|专卖|台球桌|桌球桌|维修|安装|批发')
CLOSED = re.compile(r'暂停营业|停止营业|已关闭|已歇业|永久关闭')
MUNICIPALITIES = {'110000', '120000', '310000', '500000'}


def clean(value):
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFKC', str(value or ''))).strip()


def key(value):
    return re.sub(r'\s+', '', clean(value)).casefold()


def normalize(record, nodes, province_name=None, venue_verified=False):
    d = record['data']
    name, address, phone = clean(d.get('门店名称')), clean(d.get('地址')), clean(d.get('电话'))
    poi, code, category = clean(d.get('高德 POI ID')), clean(d.get('行政区划代码')), clean(d.get('高德分类'))
    issues = []
    if not name or len(name) > 200: issues.append('INVALID_NAME')
    if not address or len(address) > 500: issues.append('MISSING_OR_INVALID_ADDRESS')
    if len(phone) > 50: issues.append('PHONE_TOO_LONG')
    if not re.fullmatch(r'B[A-Z0-9]{9}', poi): issues.append('INVALID_POI')
    if not venue_verified and SUPPLIER.search(name): issues.append('POSSIBLE_EQUIPMENT_SUPPLIER')
    if CLOSED.search(name): issues.append('POSSIBLY_CLOSED')
    if (not venue_verified and '台球厅' not in category and not (VENUE.search(name) and
            (not category or '体育休闲服务' in category))):
        issues.append('VENUE_TYPE_UNCONFIRMED')
    try:
        longitude, latitude = [float(v) for v in str(d.get('经纬度')).split(',')]
        if not (math.isfinite(longitude) and math.isfinite(latitude) and 73 <= longitude <= 136 and 3 <= latitude <= 54):
            raise ValueError()
    except (TypeError, ValueError):
        longitude, latitude = None, None
        issues.append('INVALID_COORDINATES')
    node = nodes.get(code)
    province = city = None
    if node is None:
        issues.append('ADMIN_CODE_UNMATCHED')
    else:
        chain = [node]
        while chain[-1]['parentCode']:
            chain.append(nodes[chain[-1]['parentCode']])
        province = next(n for n in chain if n['level'] == 'PROVINCE')
        city = next((n for n in chain if n['level'] == 'PREFECTURE'), None)
        has_children = any(n['parentCode'] == code for n in nodes.values()) if node['level'] != 'COUNTY' else False
        if node['level'] != 'COUNTY' and not (node['level'] == 'PREFECTURE' and not has_children):
            issues.append('ADMIN_NOT_COUNTY_OR_LEGAL_CITY_LEAF')
        if province_name and province['name'] != province_name:
            issues.append('PROVINCE_CONFLICT')
    if issues:
        return None, issues
    ref = {'file': record['file'], 'sheet': record['sheet'], 'row': record['row']}
    return {
        'name': name, 'address': address, 'publicPhone': phone or None,
        'longitude': round(longitude, 7), 'latitude': round(latitude, 7), 'coordinateSystem': 'GCJ02',
        'adminAreaCode': code, 'provinceCode': province['code'],
        'cityCode': city['code'] if city else province['code'] if province['code'] in MUNICIPALITIES else None,
        'countyCode': code if node['level'] == 'COUNTY' else None,
        'geoRegionCode': province['regionCode'],
        'areaResolution': 'COUNTY' if node['level'] == 'COUNTY' else 'LEGAL_CITY_LEAF',
        'identities': [{'poi': poi, 'source': ref}], 'sourceCategory': category,
    }, []


def deduplicate(candidates):
    """A POI may occur once; distinct POIs merge only with identical venue, address and point."""
    by_poi, by_venue, output, merged = {}, {}, [], []
    for row in candidates:
        poi = row['identities'][0]['poi']
        signature = (row['adminAreaCode'], key(row['name']), key(row['address']), row['longitude'], row['latitude'])
        if poi in by_poi:
            prior = by_poi[poi]
            if prior != signature:
                raise ValueError(f'Conflicting source POI: {poi}')
            merged.append({'kind': 'REPEATED_POI', 'poi': poi})
            continue
        by_poi[poi] = signature
        previous = by_venue.get(signature)
        # Different nonempty phone sets can indicate distinct businesses at the same address.
        if previous and previous['publicPhone'] and row['publicPhone'] and key(previous['publicPhone']) != key(row['publicPhone']):
            previous = None
        if previous:
            previous['identities'].extend(row['identities'])
            previous['publicPhone'] = previous['publicPhone'] or row['publicPhone']
            merged.append({'kind': 'EXACT_VENUE', 'keptPoi': previous['identities'][0]['poi'], 'aliasPoi': poi})
        else:
            output.append(row)
            by_venue[signature] = row
    return output, merged


def prepare(source, catalog_path, output, raw_cache=None):
    import openpyxl
    catalog = json.loads(catalog_path.read_text())
    nodes = {n['code']: n for n in catalog['nodes']}
    province_names = [n['name'] for n in nodes.values() if n['level'] == 'PROVINCE']
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    output.mkdir(parents=True, exist_ok=True)
    candidates, rejected, files, counts = [], [], set(), Counter()

    def process(record):
        if record['sheet'] not in FULL_SHEETS:
            counts['phoneSubsetRowsSkipped'] += 1
            return
        counts['candidateRows'] += 1
        files.add(record['file'])
        province = next((n for n in province_names if Path(record['file']).name.startswith(n + '_')), None)
        if province is None: raise ValueError('Cannot determine source province: ' + record['file'])
        row, issues = normalize(record, nodes, province)
        if issues:
            rejected.append({**record, 'reasons': issues})
            counts.update(issues)
        else:
            candidates.append(row)

    if raw_cache:
        # Cache comes from the exact ZIP profile; require its manifest hash before use.
        manifest = json.loads(raw_cache.with_name('profile.json').read_text())
        if manifest['zipSha256'] != digest: raise ValueError('Raw cache does not match the source ZIP')
        with raw_cache.open() as cached:
            for line in cached: process(json.loads(line))
    else:
        with ZipFile(source) as archive:
            for info in archive.infolist():
                if not info.filename.endswith('.xlsx'): continue
                if info.file_size > 50 * 1024 * 1024: raise ValueError('Unexpected workbook size')
                book = openpyxl.load_workbook(io.BytesIO(archive.read(info)), read_only=True, data_only=True)
                try:
                    if not all(s in book.sheetnames for s in FULL_SHEETS): raise ValueError('Missing full-data worksheet')
                    for sheet in book:
                        if sheet.title == '交付概览': continue
                        rows = sheet.iter_rows(values_only=True)
                        header = next(rows)
                        if not REQUIRED.issubset(header): raise ValueError('Missing required workbook columns')
                        for number, values in enumerate(rows, 2):
                            if any(v is not None for v in values):
                                process({'file': info.filename, 'sheet': sheet.title, 'row': number, 'data': dict(zip(header, values))})
                finally:
                    book.close()
    accepted, merged = deduplicate(candidates)
    accepted.sort(key=lambda row: row['identities'][0]['poi'])
    with (output/'accepted.jsonl').open('w') as target:
        for row in accepted: target.write(json.dumps(row, ensure_ascii=False) + '\n')
    with (output/'review.jsonl').open('w') as target:
        for row in rejected: target.write(json.dumps(row, ensure_ascii=False) + '\n')
    with (output/'review.csv').open('w', newline='', encoding='utf-8-sig') as target:
        writer = csv.writer(target);writer.writerow(['原因', '门店名称', 'POI', '城市', '区县', '原区划编码', '地址', '来源文件', '工作表', '行号'])
        for row in rejected:
            d=row['data']; values=[';'.join(row['reasons']),d.get('门店名称'),d.get('高德 POI ID'),d.get('城市'),d.get('区县'),d.get('行政区划代码'),d.get('地址'),row['file'],row['sheet'],row['row']]
            writer.writerow(["'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v for v in values])
    (output/'duplicates.json').write_text(json.dumps(merged,ensure_ascii=False,indent=2))
    manifest = {'sourceZipSha256': digest, 'catalogVersion': catalog['version'],
        'batchRef': 'AMAP_FILE_' + digest[:16], 'sourceFiles': len(files),
        'counts': dict(counts), 'reviewRows': len(rejected), 'duplicateRowsRemoved': len(merged),
        'acceptedStores': len(accepted), 'acceptedIdentities': sum(len(r['identities']) for r in accepted),
        'byProvince': dict(Counter(r['provinceCode'] for r in accepted)),
        'byCity': dict(Counter(r['cityCode'] or 'DIRECT_COUNTY' for r in accepted)),
        'byRegion': dict(Counter(r['geoRegionCode'] for r in accepted)),
        'acceptedSha256': hashlib.sha256((output/'accepted.jsonl').read_bytes()).hexdigest()}
    assert counts['candidateRows'] == len(rejected) + len(merged) + len(accepted)
    (output/'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2))
    print(json.dumps({k:v for k,v in manifest.items() if k not in ('byCity','byProvince')},ensure_ascii=False,indent=2))


if __name__ == '__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('source',type=Path);p.add_argument('catalog',type=Path);p.add_argument('output',type=Path)
    p.add_argument('--raw-cache',type=Path)
    a=p.parse_args();prepare(a.source,a.catalog,a.output,a.raw_cache)
