"""从同目录需求源文档生成 Word；需使用 Codex bundled Python 运行。"""
from pathlib import Path
import re
import argparse
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--source', type=Path, default=ROOT / 'COST_DASHBOARD_PRD_V4_2026-10-16.md')
parser.add_argument('--output', type=Path, default=ROOT / '成本数据看板产品方案_V4_2026-10-16.docx')
args = parser.parse_args()
SOURCE, OUTPUT = args.source, args.output
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin, sec.bottom_margin = Mm(18), Mm(17)
sec.left_margin, sec.right_margin = Mm(18), Mm(18)
sec.footer_distance = Mm(8)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    s = doc.styles[name]
    s.font.name = 'Heiti SC'
    fonts=s.element.get_or_add_rPr().rFonts
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        fonts.attrib.pop(qn('w:'+attr), None)
    fonts.set(qn('w:eastAsia'), 'Heiti SC')
    for border in s.element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
    s.font.color.rgb = RGBColor(0, 0, 0)
normal = doc.styles['Normal']
normal.font.size = Pt(10)
normal.paragraph_format.line_spacing = 1.17
normal.paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(24)
doc.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Title'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 2'].font.size = Pt(12)

footer = sec.footer.paragraphs[0]
footer.alignment = 2
r = footer.add_run('成本数据看板需求与实现方案   ')
r.font.size = Pt(8)
r.font.color.rgb = RGBColor.from_string('768699')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
footer._p.append(fld)

def text(p, s):
    # 保留源码编码，将官方调研来源生成可点击链接。
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    s = s.replace('`', '')
    last = 0
    for match in re.finditer(r'\[([^\]]+)\]\((https?://[^\s)]+)\)', s):
        p.add_run(s[last:match.start()])
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), p.part.relate_to(match.group(2), RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '31689A'); props.append(color)
        run.append(props)
        value = OxmlElement('w:t'); value.text = match.group(1); run.append(value)
        link.append(run); p._p.append(link)
        last = match.end()
    p.add_run(s[last:])

def table(lines):
    rows = [[v.strip() for v in line.strip().strip('|').split('|')] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r':?-+:?', v) for v in r)]
    n = len(rows[0]); t = doc.add_table(rows=0, cols=n)
    t.autofit = False
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        border = OxmlElement('w:'+edge); border.set(qn('w:val'),'single'); border.set(qn('w:sz'),'4'); border.set(qn('w:color'),'D9D9D9'); borders.append(border)
    t._tbl.tblPr.append(borders)
    # 源码索引含长路径，给证据列更多宽度，避免末段挤入页脚。
    widths = ([70, 104] if sections == 14 else [48, 126]) if n == 2 else [37, 61, 76] if n == 3 else [174/n]*n
    for col, width in zip(t.columns, widths): col.width = Mm(width)
    for i, cells in enumerate(rows):
        row = t.add_row()
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trPr.append(cant)
        if i == 0:
            rep = OxmlElement('w:tblHeader'); trPr.append(rep)
        for j, value in enumerate(cells):
            cell = row.cells[j]; cell.width = Mm(widths[j])
            tcPr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8EEF5' if i == 0 else ('F5F7FA' if i % 2 else 'FFFFFF')); tcPr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side, val in [('top',80),('bottom',80),('left',100),('right',100)]:
                el=OxmlElement('w:'+side); el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');margins.append(el)
            tcPr.append(margins)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.12
            text(p,value)
            for run in p.runs:
                run.font.size=Pt(9)
                run.bold=i==0
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

lines=SOURCE.read_text().splitlines();i=0;sections=0
while i < len(lines):
    line=lines[i].strip()
    if not line:
        i+=1;continue
    if line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            block.append(lines[i]);i+=1
        table(block);continue
    if line.startswith('# '):doc.add_paragraph(line[2:],style='Title')
    elif line.startswith('## '):
        p=doc.add_paragraph(line[3:],style='Heading 1')
        p.paragraph_format.page_break_before=sections>0
        sections+=1
    else:text(doc.add_paragraph(),line)
    i+=1
doc.core_properties.title='成本数据看板产品方案'
doc.core_properties.subject='现有成本台账与报销的数据看板设计'
doc.core_properties.author=''
doc.core_properties.keywords='成本台账,报销,BI,需求,数据口径'
doc.save(OUTPUT)
print(OUTPUT)
print(f'Sections: {sections}')
